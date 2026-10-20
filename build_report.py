import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_page_number_field(run):
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def style_paragraph(p, font_name="Times New Roman", size_pt=12, bold=False, italic=False, color_rgb=(0,0,0), align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, space_before=0, line_spacing=1.5):
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    for run in p.runs:
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = RGBColor(*color_rgb)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.add_run(text)
    style_paragraph(p, size_pt=16, bold=True, color_rgb=(15, 23, 42), align=WD_ALIGN_PARAGRAPH.CENTER, space_before=16, space_after=10, line_spacing=1.2)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.add_run(text)
    style_paragraph(p, size_pt=14, bold=True, color_rgb=(30, 41, 59), align=WD_ALIGN_PARAGRAPH.LEFT, space_before=12, space_after=6, line_spacing=1.2)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.add_run(text)
    style_paragraph(p, size_pt=12, bold=True, color_rgb=(51, 65, 85), align=WD_ALIGN_PARAGRAPH.LEFT, space_before=8, space_after=4, line_spacing=1.2)
    return p

def add_body_p(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    p.add_run(text)
    style_paragraph(p, size_pt=12, bold=bold, italic=italic, space_after=6, line_spacing=1.5)
    return p

def add_bullet_p(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    r1 = p.add_run(bold_prefix)
    r1.bold = True
    r2 = p.add_run(text)
    style_paragraph(p, size_pt=12, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4, line_spacing=1.3)
    return p

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "0F172A")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    set_table_borders(tbl, color="334155", sz="6")
    p = cell.paragraphs[0]
    p.add_run(code_text)
    style_paragraph(p, font_name="Courier New", size_pt=9.5, color_rgb=(241, 245, 249), align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0, line_spacing=1.15)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_custom_table(table, col_widths, headers, rows_data):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="CBD5E1", sz="4")
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        set_cell_background(hdr_cells[i], "1E293B")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        style_paragraph(p, size_pt=10, bold=True, color_rgb=(255,255,255), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0, line_spacing=1.1)

    # Data Rows
    for row_idx, row_data in enumerate(rows_data):
        row_cells = table.rows[row_idx + 1].cells
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = str(cell_value)
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=100, bottom=100, left=140, right=140)
            p = row_cells[col_idx].paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.CENTER if (col_idx in [0, 2, 3] and len(headers) > 3) else WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(p, size_pt=10, color_rgb=(30,41,59), align=align, space_after=0, line_spacing=1.2)

    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

print("Base setup ready.")
